#!/usr/bin/env python3
"""Normalize Office inputs to cached Markdown files."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {".md", ".markdown", ".docx", ".xlsx"}
INVALID_PATH_CHARS_RE = re.compile(r'[<>:"/\\|?*\x00-\x1f]+')


def repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_stem(path: Path) -> str:
    stem = INVALID_PATH_CHARS_RE.sub("_", path.stem).strip(" .")
    return stem or "document"


def escape_table_cell(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    text = text.replace("|", r"\|")
    return text.replace("\n", "<br>")


def markdown_table(rows: list[list[str]]) -> list[str]:
    if not rows:
        return []
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    header = [cell if cell else f"列{index + 1}" for index, cell in enumerate(padded[0])]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in padded[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return lines


def docx_image_count(doc: Any) -> int:
    try:
        inline_count = len(doc.inline_shapes)
    except Exception:
        inline_count = 0
    rel_count = 0
    try:
        rel_count = sum(1 for rel in doc.part.rels.values() if "image" in rel.reltype)
    except Exception:
        rel_count = 0
    return max(inline_count, rel_count)


def iter_docx_blocks(parent: Any):
    from docx.document import Document as DocxDocument
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent_element = parent.element.body if isinstance(parent, DocxDocument) else parent._tc
    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def paragraph_to_markdown(paragraph: Any) -> list[str]:
    text = paragraph.text.strip()
    if not text:
        return []

    style_name = ""
    try:
        style_name = (paragraph.style.name or "").strip()
    except Exception:
        style_name = ""
    style_lower = style_name.lower()

    if style_lower.startswith("toc") or "table of contents" in style_lower:
        return []

    heading_match = re.search(r"heading\s*(\d+)", style_lower)
    if not heading_match:
        heading_match = re.search(r"标题\s*(\d+)", style_lower)
    if heading_match:
        level = min(max(int(heading_match.group(1)), 1), 6)
        return [f"{'#' * level} {text}"]

    if "list" in style_lower or "列表" in style_lower:
        return [f"- {text}"]

    return [text]


def table_to_markdown(table: Any) -> list[str]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [escape_table_cell(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)
    return markdown_table(rows)


def convert_docx(source: Path, output: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for .docx conversion") from exc

    doc = Document(str(source))
    image_count = docx_image_count(doc)
    lines: list[str] = [f"# {source.stem}", ""]

    for block in iter_docx_blocks(doc):
        if block.__class__.__name__ == "Paragraph":
            converted = paragraph_to_markdown(block)
        else:
            converted = table_to_markdown(block)
        if converted:
            lines.extend(converted)
            lines.append("")

    output.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")

    warnings: list[str] = []
    if image_count:
        warnings.append(
            f"Detected {image_count} image(s). Text/table conversion completed, "
            "but diagrams or screenshots require the local DOCX image and diagram supplement workflow."
        )
    return {"kind": "docx", "image_count": image_count, "warnings": warnings}


def convert_xlsx(source: Path, output: Path) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for .xlsx conversion") from exc

    workbook = openpyxl.load_workbook(source, data_only=True)
    lines: list[str] = [f"# {source.stem}", ""]
    sheet_count = 0
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [escape_table_cell(value) for value in row]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        sheet_count += 1
        lines.extend([f"## {sheet.title}", ""])
        lines.extend(markdown_table(rows))
        lines.append("")

    output.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return {"kind": "xlsx", "sheet_count": sheet_count, "warnings": []}


def normalize_one(source: Path, cache_dir: Path, force: bool) -> dict[str, Any]:
    source = source.expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(f"Input file does not exist: {source}")
    if not source.is_file():
        raise ValueError(f"Input path is not a file: {source}")

    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported input type {suffix}: {source}")

    if suffix in {".md", ".markdown"}:
        return {
            "source": str(source),
            "markdown": str(source),
            "cached": False,
            "kind": "markdown",
            "warnings": [],
        }

    sha = file_sha256(source)
    out_dir = cache_dir / sha[:12]
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = safe_stem(source)
    markdown_path = out_dir / f"{stem}.md"
    metadata_path = out_dir / f"{stem}.conversion.json"

    cached = markdown_path.exists() and metadata_path.exists() and not force
    conversion_extra: dict[str, Any] = {"kind": suffix.lstrip("."), "warnings": []}
    if not cached:
        if suffix == ".docx":
            conversion_extra = convert_docx(source, markdown_path)
        elif suffix == ".xlsx":
            conversion_extra = convert_xlsx(source, markdown_path)

        stat = source.stat()
        metadata = {
            "source": str(source),
            "source_name": source.name,
            "source_size": stat.st_size,
            "source_mtime": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
            "sha256": sha,
            "markdown": str(markdown_path.resolve()),
            "converted_at": datetime.now(timezone.utc).isoformat(),
            **conversion_extra,
        }
        metadata_path.write_text(
            json.dumps(metadata, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    else:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        conversion_extra = {
            "kind": metadata.get("kind", suffix.lstrip(".")),
            "warnings": metadata.get("warnings", []),
        }

    return {
        "source": str(source),
        "markdown": str(markdown_path.resolve()),
        "metadata": str(metadata_path.resolve()),
        "cached": cached,
        "sha256": sha,
        **conversion_extra,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Normalize .docx/.xlsx inputs to cached Markdown")
    parser.add_argument("inputs", nargs="+", type=Path, help="Input .md, .docx or .xlsx files")
    parser.add_argument(
        "--cache-dir",
        type=Path,
        default=repo_root() / "outputs" / "input-cache",
        help="Cache directory, defaults to outputs/input-cache",
    )
    parser.add_argument("--force", action="store_true", help="Regenerate cached Markdown")
    parser.add_argument("--json", action="store_true", help="Print machine-readable JSON")
    args = parser.parse_args()

    cache_dir = args.cache_dir
    if not cache_dir.is_absolute():
        cache_dir = repo_root() / cache_dir
    cache_dir.mkdir(parents=True, exist_ok=True)

    results = []
    try:
        for source in args.inputs:
            results.append(normalize_one(source, cache_dir, args.force))
    except Exception as exc:
        print(f"失败: {exc}", file=sys.stderr)
        return 1

    if args.json:
        print(json.dumps(results, ensure_ascii=False, indent=2))
    else:
        for result in results:
            if result["kind"] == "markdown":
                print(f"无需转换: {result['source']}")
                continue
            state = "复用缓存" if result["cached"] else "已转换"
            print(f"{state}: {result['source']} -> {result['markdown']}")
            for warning in result.get("warnings", []):
                print(f"警告: {warning}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
